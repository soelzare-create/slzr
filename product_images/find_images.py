"""Find product photos on Digikala by barcode / product name.

Run this on a machine inside Iran (Digikala blocks most foreign IPs):

    pip install requests openpyxl
    python find_images.py products.xlsx

Output:
    images/<barcode>.jpg   downloaded photos
    result.xlsx            original sheet + found title, similarity, page URL, image URL
    cache.json             resume cache (re-running skips rows already done)
"""
import difflib
import json
import os
import re
import sys
import time

import openpyxl
import requests

SEARCH = "https://api.digikala.com/v1/search/"
HEADERS = {"User-Agent": "Mozilla/5.0", "Accept": "application/json"}
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
MIN_SIMILARITY = 0.45  # below this, the match is flagged for manual review


def norm(text):
    text = (text or "").translate(PERSIAN_DIGITS).replace("ي", "ی").replace("ك", "ک")
    return re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip().lower()


def search(query):
    r = requests.get(SEARCH, params={"q": query, "page": 1}, headers=HEADERS, timeout=20)
    r.raise_for_status()
    return (r.json().get("data") or {}).get("products") or []


def first_image(product):
    urls = ((product.get("images") or {}).get("main") or {}).get("url") or []
    return urls[0].split("?")[0] if urls else None


def best_match(name, products):
    best, best_score = None, 0.0
    for p in products[:10]:
        score = difflib.SequenceMatcher(None, norm(name), norm(p.get("title_fa"))).ratio()
        if score > best_score and first_image(p):
            best, best_score = p, score
    return best, best_score


def lookup(name, barcode):
    # 1) barcode search: any hit is trusted; 2) name search: best fuzzy title match.
    for query, by_barcode in ((barcode, True), (name, False)):
        if not query:
            continue
        products = search(query)
        if not products:
            continue
        p, score = best_match(name, products)
        if p and (by_barcode or score >= 0.25):
            return {
                "via": "barcode" if by_barcode else "name",
                "title": p.get("title_fa"),
                "score": round(score, 2),
                "page": "https://www.digikala.com/product/dkp-%s/" % p.get("id"),
                "image": first_image(p),
            }
    return None


def main(path):
    os.makedirs("images", exist_ok=True)
    cache = json.load(open("cache.json", encoding="utf-8")) if os.path.exists("cache.json") else {}
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    extra = ["منبع", "روش یافتن", "عنوان یافت‌شده", "شباهت نام", "نیاز به بررسی", "صفحهٔ محصول", "فایل عکس"]
    base = ws.max_column
    for i, h in enumerate(extra, 1):
        ws.cell(1, base + i, h)

    rows = list(ws.iter_rows(min_row=2))
    for n, row in enumerate(rows, 1):
        name, barcode = str(row[1].value or ""), str(row[2].value or "").strip()
        key = barcode or name
        if key not in cache:
            try:
                cache[key] = lookup(name, barcode)
            except Exception as e:  # network / rate limit: leave for the next run
                print("  error:", e)
                time.sleep(5)
                continue
            json.dump(cache, open("cache.json", "w", encoding="utf-8"), ensure_ascii=False)
            time.sleep(1)
        hit = cache[key]
        print(f"[{n}/{len(rows)}] {name} -> {hit and hit['title']}")
        if not hit:
            continue
        fname = os.path.join("images", re.sub(r"[^\w-]", "_", key) + ".jpg")
        if not os.path.exists(fname):
            try:
                img = requests.get(hit["image"], headers=HEADERS, timeout=30)
                img.raise_for_status()
                open(fname, "wb").write(img.content)
            except Exception as e:
                print("  image error:", e)
                fname = ""
        r = row[0].row
        review = "بله" if hit["score"] < MIN_SIMILARITY else ""
        for i, v in enumerate(["دیجی‌کالا", hit["via"], hit["title"], hit["score"], review, hit["page"], fname], 1):
            ws.cell(r, base + i, v)
        ws.cell(r, 8, hit["image"])  # «آدرس عکس اصلی» column
    wb.save("result.xlsx")
    found = sum(1 for v in cache.values() if v)
    print(f"done: {found}/{len(cache)} found -> result.xlsx, images/")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "products.xlsx")
