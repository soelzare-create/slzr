#!/usr/bin/env python3
"""Barcode list (Excel) -> product photos -> RTL Excel with embedded images.

Script version of the "shelf-photo-to-planogram-excel" skill. Instead of reading a
shelf photo, it takes the barcodes (and names) from an Excel file, looks up a real
packshot for each one on Iranian store sites, and writes an output Excel with the
image embedded in each row.

Lookup order (same sources as the skill, barcode-exact sources first):
  1. daychi.com      /products/<EAN>  -> the URL *is* the barcode       (exact)
  2. hamekara.com    /api/v3/search?q=<EAN> -> JSON with "barcode"      (exact)
  3. rojashop.com    product page text "بارکد : <EAN>"                 (exact)
  4. openfoodfacts   /api/v2/product/<EAN>.json                         (exact)
  5. digikala.com    search API by product name                         (name match)
  6. torob.com       search API by product name                         (name match)

Rules carried over from the skill:
  * never invent an image or barcode; unknown -> «پیدا نشد»
  * name-search results are accepted only if the name and the size/weight match
  * polite pacing (delay between requests); a source that answers 403/429 or a
    captcha is switched off for the rest of the run (never bypassed)
  * every output file is RTL

Iranian sites are not reachable from foreign servers: run this on a computer in Iran.

usage:
  python find_product_photos.py products.xlsx
  python find_product_photos.py products.xlsx -o out.xlsx --cutout --limit 50
  python find_product_photos.py --help
"""
from __future__ import annotations

import argparse
import csv
import html
import io
import json
import os
import random
import re
import shutil
import sys
import tempfile
import time
from dataclasses import dataclass, field, asdict
from urllib.parse import quote, urljoin

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")

HIGH, MEDIUM, LOW, NOT_FOUND = "بالا", "متوسط", "پایین", "پیدا نشد"

# --------------------------------------------------------------------------- text

_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
_CHAR_FIX = str.maketrans({"ي": "ی", "ى": "ی", "ك": "ک", "ة": "ه", "ۀ": "ه", "أ": "ا",
                           "إ": "ا", "آ": "ا", "ؤ": "و", "‌": " ", "‏": "",
                           "‎": "", "ـ": ""})
_STOP = {"و", "با", "در", "از", "به", "مدل", "حجم", "وزن", "مقدار", "بسته", "ی", "ای",
         "عدد", "عددی", "گرم", "گرمی", "میلی", "لیتر", "لیتری", "کیلو", "کیلوگرم",
         "کیلویی", "سی", "گرمی", "بطری", "شیشه", "قوطی", "پت"}


def to_latin_digits(s: str) -> str:
    return (s or "").translate(_FA_DIGITS)


def normalize(s: str) -> str:
    s = to_latin_digits(str(s or "")).translate(_CHAR_FIX).lower()
    s = re.sub(r"(\d)([^\d\s.])", r"\1 \2", s)          # "1لیتر" -> "1 لیتر"
    s = re.sub(r"([^\d\s.])(\d)", r"\1 \2", s)
    s = re.sub(r"[^\w\s.]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def tokens(s: str) -> set[str]:
    return {t for t in normalize(s).split() if t not in _STOP and not t.replace(".", "").isdigit()
            and len(t) > 1}


def name_score(wanted: str, candidate: str) -> float:
    """Share of the wanted name's words that appear in the candidate title (0..1)."""
    a, b = tokens(wanted), tokens(candidate)
    if not a or not b:
        return 0.0
    hit = 0.0
    for t in a:
        if t in b:
            hit += 1
        elif len(t) >= 4 and any(t in u or u in t for u in b if len(u) >= 4):
            hit += 0.7                                   # «ابمیوه» vs «آبمیوه» etc.
    return hit / len(a)


_UNITS = [  # (regex, kind, factor to base unit)
    (r"کیلو\s*گرم|کیلوگرمی|کیلوگرم|کیلویی|کیلو|kg", "w", 1000),
    (r"گرمی|گرم|gr|g", "w", 1),
    (r"میلی\s*لیتری|میلی\s*لیتر|ml|سی\s*سی|cc", "v", 1),
    (r"لیتری|لیتر|l", "v", 1000),
    (r"عددی|عدد|تایی|رول", "n", 1),
]


def sizes(s: str) -> dict[str, set[float]]:
    """Extract sizes: {'w': {480.0}, 'v': {...}, 'n': {...}} (grams / ml / count)."""
    s = normalize(s).replace("/", ".")
    out: dict[str, set[float]] = {}
    for unit_re, kind, f in _UNITS:
        for m in re.finditer(r"(\d+(?:\.\d+)?)\s*(?:" + unit_re + r")(?![\w])", s):
            out.setdefault(kind, set()).add(round(float(m.group(1)) * f, 2))
        s = re.sub(r"(\d+(?:\.\d+)?)\s*(?:" + unit_re + r")(?![\w])", " ", s)
    return out


def size_match(wanted: str, candidate: str) -> bool | None:
    """True = same size, False = different size, None = can't tell."""
    a, b = sizes(wanted), sizes(candidate)
    verdict = None
    for kind in ("w", "v", "n"):
        if kind in a and kind in b:
            if a[kind] & b[kind]:
                verdict = True if verdict is None else verdict
            else:
                return False
    return verdict


def ean_ok(code: str) -> bool:
    if not re.fullmatch(r"\d{13}", code or ""):
        return False
    s = sum(int(c) * (1 if i % 2 == 0 else 3) for i, c in enumerate(code[:12]))
    return (10 - s % 10) % 10 == int(code[12])


def clean_barcode(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\D", "", to_latin_digits(str(v)))


def safe_filename(s: str, maxlen: int = 120) -> str:
    s = re.sub(r'[\\/:*?"<>|\r\n\t]', " ", str(s))
    return re.sub(r"\s+", " ", s).strip()[:maxlen].strip() or "item"


# --------------------------------------------------------------------------- http

class SourceBlocked(Exception):
    """The site answered 403/429/captcha: stop using it for this run."""


class Http:
    def __init__(self, delay: float, timeout: float = 20, verbose: bool = False):
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": UA, "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.5"})
        self.delay, self.timeout, self.verbose = delay, timeout, verbose
        self._last = 0.0

    def _wait(self):
        gap = self.delay * random.uniform(1.0, 2.0) - (time.time() - self._last)
        if gap > 0:
            time.sleep(gap)
        self._last = time.time()

    def get(self, url: str, **kw) -> requests.Response | None:
        for attempt in range(3):
            self._wait()
            try:
                r = self.s.get(url, timeout=self.timeout, **kw)
            except requests.RequestException as e:
                if self.verbose:
                    print(f"    ! {url[:90]} {type(e).__name__}", flush=True)
                time.sleep(2 * (attempt + 1))
                continue
            if r.status_code in (403, 429):
                raise SourceBlocked(f"HTTP {r.status_code}")
            if r.status_code >= 500:
                time.sleep(2 * (attempt + 1))
                continue
            if "text/html" in r.headers.get("content-type", "") and re.search(
                    r"captcha|recaptcha|cf-chl|arvancloud.*challenge", r.text[:20000], re.I):
                raise SourceBlocked("captcha")
            return r
        return None

    def json(self, url: str, **kw):
        r = self.get(url, headers={"Accept": "application/json"}, **kw)
        if r is None or r.status_code != 200:
            return None
        try:
            return r.json()
        except ValueError:
            return None


# --------------------------------------------------------------------------- sources

@dataclass
class Candidate:
    source: str
    title: str
    image_url: str
    page_url: str = ""
    barcode_exact: bool = False   # the source itself ties this item to our barcode


def _find_images(obj, base: str = "") -> list[str]:
    """Collect image-looking URLs from an arbitrary JSON object (unknown schemas)."""
    found: list[str] = []

    def walk(o, key=""):
        if isinstance(o, dict):
            for k, v in o.items():
                walk(v, k.lower())
        elif isinstance(o, list):
            for v in o:
                walk(v, key)
        elif isinstance(o, str):
            if re.search(r"\.(jpe?g|png|webp)(\?|$)", o, re.I) or ("image" in key and "/" in o):
                found.append(urljoin(base, o) if base else o)

    walk(obj)
    return found


def src_daychi(http: Http, code: str, name: str) -> list[Candidate]:
    base = "https://www.daychi.com"
    r = http.get(f"{base}/products/{code}")
    if r is None or r.status_code != 200 or code not in r.url:
        return []
    h = r.text
    title = ""
    for pat in (r'<meta[^>]+property=["\']og:title["\'][^>]+content=["\']([^"\']+)',
                r"<h1[^>]*>(.*?)</h1>", r"<title>(.*?)</title>"):
        m = re.search(pat, h, re.S | re.I)
        if m:
            title = html.unescape(re.sub(r"<[^>]+>", " ", m.group(1))).strip()
            break
    imgs = re.findall(r"https://www\.daychi\.com/storage/media/[^\"'\s)]+?\.(?:jpe?g|png|webp|avif)", h)
    og = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)', h, re.I)
    if og:
        imgs.insert(0, html.unescape(og.group(1)))
    imgs = sorted(dict.fromkeys(imgs), key=lambda u: u.lower().endswith(".avif"))  # avif last
    if not title or not imgs:
        return []
    return [Candidate("daychi", title, imgs[0], r.url, barcode_exact=True)]


def src_hamekara(http: Http, code: str, name: str) -> list[Candidate]:
    base = "https://hamekara.com"
    j = http.json(f"{base}/api/v3/search?q={quote(code)}")
    ps = (j or {}).get("data", {}).get("products", []) if isinstance(j, dict) else []
    out = []
    for p in ps:
        if clean_barcode(p.get("barcode")) != code:
            continue
        imgs = _find_images(p, base)
        if imgs:
            slug = p.get("slug") or p.get("id") or ""
            out.append(Candidate("hamekara", str(p.get("title", "")).strip(), imgs[0],
                                 f"{base}/product/{slug}" if slug else "", barcode_exact=True))
    return out


def src_rojashop(http: Http, code: str, name: str, max_pages: int = 3) -> list[Candidate]:
    base = "https://www.rojashop.com"
    r = http.get(f"{base}/search?term={quote(code)}")
    if r is None or r.status_code != 200:
        return []
    links = re.findall(r'<a[^>]+class=["\'][^"\']*productLink[^"\']*["\'][^>]*href=["\']([^"\']+)', r.text)
    links += re.findall(r'<a[^>]+href=["\']([^"\']+)["\'][^>]*class=["\'][^"\']*productLink', r.text)
    out = []
    for u in list(dict.fromkeys(urljoin(base, x.split("?")[0]) for x in links))[:max_pages]:
        r2 = http.get(u)
        if r2 is None or r2.status_code != 200:
            continue
        text = re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", r2.text)))
        bar = re.search(r"بارکد\s*:?\s*(\d{8,14})", to_latin_digits(text))
        og = re.search(r'<meta[^>]+property=["\']og:image["\'][^>]+content=["\']([^"\']+)', r2.text, re.I)
        t = re.search(r'<meta[^>]+property=["\']og:title["\'][^>]+content=["\']([^"\']+)', r2.text, re.I)
        if bar and bar.group(1) == code and og:
            out.append(Candidate("rojashop", html.unescape(t.group(1)) if t else "", og.group(1), u,
                                 barcode_exact=True))
            break
    return out


def src_off(http: Http, code: str, name: str) -> list[Candidate]:
    j = http.json(f"https://world.openfoodfacts.org/api/v2/product/{code}.json"
                  "?fields=product_name,product_name_fa,brands,image_front_url,image_url")
    if not j or j.get("status") != 1:
        return []
    p = j.get("product", {})
    img = p.get("image_front_url") or p.get("image_url")
    if not img:
        return []
    title = " ".join(x for x in (p.get("product_name_fa") or p.get("product_name"), p.get("brands")) if x)
    return [Candidate("openfoodfacts", title, img, f"https://world.openfoodfacts.org/product/{code}",
                      barcode_exact=True)]


def src_digikala(http: Http, code: str, name: str) -> list[Candidate]:
    out = []
    for q in (name,):
        j = http.json(f"https://api.digikala.com/v1/search/?q={quote(q)}&page=1")
        ps = ((j or {}).get("data") or {}).get("products") or []
        for p in ps[:20]:
            urls = (((p.get("images") or {}).get("main") or {}).get("url")) or []
            if not urls:
                continue
            img = urls[0].split("?")[0]
            out.append(Candidate("digikala", p.get("title_fa", ""), img,
                                 f"https://www.digikala.com/product/dkp-{p.get('id')}/"))
    return out


def src_torob(http: Http, code: str, name: str) -> list[Candidate]:
    j = http.json("https://api.torob.com/v4/base-product/search/?q="
                  f"{quote(name)}&page=0&size=10&sort=popularity")
    out = []
    for p in (j or {}).get("results", []) if isinstance(j, dict) else []:
        if p.get("image_url"):
            rk = p.get("random_key", "")
            out.append(Candidate("torob", p.get("name1", ""), p["image_url"],
                                 f"https://torob.com/p/{rk}/" if rk else ""))
    return out


SOURCES = {
    "daychi": src_daychi,
    "hamekara": src_hamekara,
    "rojashop": src_rojashop,
    "off": src_off,
    "digikala": src_digikala,
    "torob": src_torob,
}
DEFAULT_SOURCES = "daychi,hamekara,rojashop,off,digikala,torob"


# --------------------------------------------------------------------------- matching

def judge(c: Candidate, name: str, min_score: float) -> tuple[str | None, str]:
    """Return (confidence or None if rejected, note)."""
    sc = name_score(name, c.title) if c.title else 0.0
    sm = size_match(name, c.title) if c.title else None
    note = f"شباهت نام {sc:.0%}"
    if sm is False:
        note += "، اندازه/وزن متفاوت"
    if c.barcode_exact:
        if sm is False:
            return LOW, "بارکد یکسان ولی " + note.split("، ")[-1]
        if sc >= 0.4 or not c.title:
            return HIGH, "بارکد دقیق در منبع؛ " + note
        return MEDIUM, "بارکد دقیق در منبع ولی نام متفاوت؛ " + note
    # name-only sources: must look like the same product
    if sm is False or sc < min_score:
        return None, note
    if sc >= 0.8 and sm is True:
        return MEDIUM, "تطبیق با نام (بدون بارکد)؛ " + note
    return LOW, "تطبیق با نام (بدون بارکد)؛ " + note


_RANK = {HIGH: 3, MEDIUM: 2, LOW: 1}


# --------------------------------------------------------------------------- images

def download_image(http: Http, url: str, dest_base: str) -> str | None:
    """Download, validate with PIL, save as .jpg/.png. Return path or None."""
    r = http.get(url, headers={"Accept": "image/avif,image/webp,image/*,*/*;q=0.8"})
    if r is None or r.status_code != 200 or len(r.content) < 3000:
        return None
    try:
        im = Image.open(io.BytesIO(r.content))
        im.load()
    except Exception:
        try:  # optional AVIF support for older Pillow
            import pillow_avif  # noqa: F401
            im = Image.open(io.BytesIO(r.content))
            im.load()
        except Exception:
            return None
    if min(im.size) < 120:
        return None
    has_alpha = im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info)
    path = dest_base + (".png" if has_alpha else ".jpg")
    if has_alpha:
        im.convert("RGBA").save(path, optimize=True)
    else:
        im.convert("RGB").save(path, quality=92)
    return path


class Cutter:
    """Background removal, same method as the skill's cutout.py (needs rembg)."""

    def __init__(self, model: str):
        from rembg import new_session, remove  # noqa: F401
        import numpy as np
        from scipy import ndimage
        self.np, self.ndi, self.remove = np, ndimage, remove
        self.session = new_session(model)

    def _keep_main(self, alpha, margin_frac=0.02):
        np, ndi = self.np, self.ndi
        mask = alpha > 60
        lab, n = ndi.label(mask)
        if n <= 1:
            return alpha
        idx = range(1, n + 1)
        area = ndi.sum(mask, lab, idx)
        H, W = mask.shape
        best, main = -1.0, 1
        for i, (cy, cx) in enumerate(ndi.center_of_mass(mask, lab, idx)):
            d = (((cy - H / 2) / H) ** 2 + ((cx - W / 2) / W) ** 2) ** 0.5
            s = area[i] * np.exp(-(d / 0.25) ** 2)
            if s > best:
                best, main = s, i + 1
        near = ndi.binary_dilation(lab == main, iterations=max(3, int(max(H, W) * margin_frac)))
        keep = np.zeros(n + 1, bool)
        keep[main] = True
        for i in np.unique(lab[near & mask]):
            if i > 0 and area[i - 1] > 0.001 * area[main - 1]:
                keep[i] = True
        out = alpha.copy()
        out[~keep[lab]] = 0
        return out

    def __call__(self, src: str, dst: str) -> str | None:
        from PIL import ImageFilter
        np = self.np
        im = Image.open(src).convert("RGB")
        if max(im.size) < 900:
            k = 900 / max(im.size)
            im = im.resize((int(im.width * k), int(im.height * k)), Image.LANCZOS)
        arr = np.array(self.remove(im, session=self.session, post_process_mask=True).convert("RGBA"))
        alpha = np.array(Image.fromarray(self._keep_main(arr[:, :, 3])).filter(ImageFilter.GaussianBlur(0.6)))
        arr[:, :, 3] = alpha
        ys, xs = np.where(alpha > 10)
        if len(ys) == 0:
            return None
        h, w = alpha.shape
        mg = int(0.02 * max(h, w))
        res = Image.fromarray(arr[max(0, ys.min() - mg):min(h, ys.max() + mg + 1),
                                  max(0, xs.min() - mg):min(w, xs.max() + mg + 1)])
        res.save(dst, optimize=True)
        return dst


# --------------------------------------------------------------------------- core

@dataclass
class Result:
    barcode: str
    status: str = NOT_FOUND
    source: str = ""
    source_title: str = ""
    image_url: str = ""
    page_url: str = ""
    image_file: str = ""
    note: str = ""
    sources_tried: list = field(default_factory=list)


def lookup(http: Http, code: str, name: str, sources: list[str], blocked: set[str],
           img_dir: str, min_score: float, verbose: bool) -> Result:
    res = Result(barcode=code)
    best: tuple[int, Candidate, str] | None = None
    if not code:
        res.note = "بارکد خالی است"
    elif not ean_ok(code) and len(code) == 13:
        res.note = "رقم کنترل EAN-13 نامعتبر است"
    for sname in sources:
        if sname in blocked:
            continue
        if not code and sname in ("daychi", "hamekara", "rojashop", "off"):
            continue
        if not name and sname in ("digikala", "torob"):
            continue
        res.sources_tried.append(sname)
        try:
            cands = SOURCES[sname](http, code, name)
        except SourceBlocked as e:
            blocked.add(sname)
            print(f"  ⚠ {sname} مسدود شد ({e}) — بقیهٔ اجرا بدون این منبع", flush=True)
            continue
        except Exception as e:  # a broken parser must not stop the whole run
            if verbose:
                print(f"  ! {sname}: {type(e).__name__}: {e}", flush=True)
            continue
        for c in cands:
            conf, note = judge(c, name, min_score)
            if conf is None:
                continue
            rank = _RANK[conf] + (0.5 if c.barcode_exact else 0) + name_score(name, c.title) / 10
            if best is None or rank > best[0]:
                best = (rank, c, (conf, note))
        if best and best[2][0] == HIGH:
            break  # good enough, stop asking other sites
    if best is None:
        return res
    # try to download, falling back through the remaining candidates is overkill: one try
    _, c, (conf, note) = best
    dest = os.path.join(img_dir, safe_filename(f"{code or 'no-code'} - {name}"))
    try:
        path = download_image(http, c.image_url, dest)
    except SourceBlocked:
        path = None
    if not path:
        res.note = (res.note + "؛ " if res.note else "") + f"عکس از {c.source} دانلود نشد"
        res.source, res.source_title, res.page_url, res.image_url = c.source, c.title, c.page_url, c.image_url
        return res
    res.status, res.source, res.source_title = conf, c.source, c.title
    res.image_url, res.page_url, res.image_file = c.image_url, c.page_url, path
    res.note = (res.note + "؛ " if res.note else "") + note
    return res


def find_col(headers: list[str], wanted: str | None, fallbacks: list[str]) -> int | None:
    norm = [normalize(h) for h in headers]
    for w in ([wanted] if wanted else []) + fallbacks:
        if w and normalize(w) in norm:
            return norm.index(normalize(w))
    return None


def read_input(path: str, sheet: str | None, barcode_col: str | None, name_col: str | None):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        sys.exit("فایل ورودی خالی است")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    bi = find_col(headers, barcode_col, ["بارکد", "barcode", "ean", "کد بارکد"])
    ni = find_col(headers, name_col, ["نام محصول", "نام کالا", "نام", "name", "title", "product"])
    ci = find_col(headers, None, ["کد کالا"])
    if bi is None:
        sys.exit(f"ستون بارکد پیدا نشد. ستون‌ها: {headers} — با --barcode-col مشخص کنید")
    data = [list(r) + [None] * (len(headers) - len(r)) for r in rows[1:] if any(v is not None for v in r)]
    return headers, data, bi, ni, ci


def build_excel(out: str, headers: list[str], data: list[list], results: list[Result | None],
                bi: int, img_col_name: str | None, thumb_h: int, summary: dict):
    extra = ["وضعیت عکس", "اطمینان", "منبع", "عنوان در منبع", "لینک محصول", "یادداشت"]
    url_i = find_col(headers, img_col_name, ["آدرس عکس اصلی", "آدرس عکس", "image url", "image"])
    defect_i = find_col(headers, None, ["نوع نقص"])
    cols = headers + extra + ["تصویر کالا"]
    img_col = len(cols)
    img_letter = get_column_letter(img_col)

    wb = Workbook()
    ws = wb.active
    ws.title = "محصولات و عکس"
    ws.sheet_view.rightToLeft = True
    ws.append(cols)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {HIGH: "C6EFCE", MEDIUM: "FFF2CC", LOW: "FCE4D6", NOT_FOUND: "F8CBAD"}
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = {"نام محصول": 42, "نام کالا": 42, "بارکد": 19, "کد کالا": 19, "شناسهٔ محصول": 14,
              "آدرس عکس اصلی": 30, "عنوان در منبع": 40, "یادداشت": 34, "لینک محصول": 26,
              "اطمینان": 11, "منبع": 12, "وضعیت عکس": 12, "نوع نقص": 14}
    for i, h in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
    ws.column_dimensions[img_letter].width = thumb_h / 7 + 4

    code_i = find_col(headers, None, ["کد کالا"])
    text_cols = [bi + 1] + ([code_i + 1] if code_i is not None and code_i != bi else [])
    tmp = tempfile.mkdtemp(prefix="thumbs_")
    n_img = 0
    for n, (row, res) in enumerate(zip(data, results), start=2):
        row = list(row)
        res = res or Result(barcode=clean_barcode(row[bi]), note="بررسی نشد (--limit)")
        found = bool(res.image_file and os.path.exists(res.image_file))
        if url_i is not None and found and not row[url_i]:
            row[url_i] = res.image_url
        if defect_i is not None and found:
            row[defect_i] = "عکس پیدا شد (نیاز به تأیید)" if res.status != HIGH else "عکس پیدا شد"
        vals = row + ["دارد" if found else "ندارد", res.status if found else NOT_FOUND,
                      res.source, res.source_title, res.page_url, res.note]
        ws.append(vals)
        for i in range(1, img_col + 1):
            c = ws.cell(n, i)
            c.border = border
            c.alignment = Alignment(vertical="center", wrap_text=True,
                                    horizontal="center" if cols[i - 1] not in widths or widths[cols[i - 1]] < 20 else "right")
        for i in text_cols:
            c = ws.cell(n, i)
            c.value = None if c.value is None else clean_barcode(c.value) or str(c.value)
            c.number_format = "@"
            c.font = Font(name="Consolas", size=11, bold=True)
        conf_cell = ws.cell(n, len(headers) + 2)
        if conf_cell.value in fills:
            conf_cell.fill = PatternFill("solid", fgColor=fills[conf_cell.value])
        link = ws.cell(n, len(headers) + 5)
        if link.value:
            link.hyperlink = link.value
            link.font = Font(color="0563C1", underline="single")
        if found:
            im = Image.open(res.image_file)
            k = thumb_h / im.height
            im = im.resize((max(1, int(im.width * k)), thumb_h), Image.LANCZOS)
            if im.width > thumb_h * 1.6:  # very wide pictures: cap width too
                k2 = thumb_h * 1.6 / im.width
                im = im.resize((int(im.width * k2), max(1, int(im.height * k2))), Image.LANCZOS)
            if im.mode in ("RGBA", "LA", "P"):
                tp = os.path.join(tmp, f"{n}.png")
                im.convert("RGBA").save(tp, optimize=True)
            else:
                tp = os.path.join(tmp, f"{n}.jpg")
                im.convert("RGB").save(tp, quality=85)
            xi = XLImage(tp)
            xi.anchor = f"{img_letter}{n}"
            ws.add_image(xi)
            ws.row_dimensions[n].height = thumb_h * 0.75 + 6
            n_img += 1
        else:
            c = ws.cell(n, img_col, "تصویر موجود نیست")
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[n].height = 30
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(img_col - 1)}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    g = wb.create_sheet("راهنما")
    g.sheet_view.rightToLeft = True
    g.column_dimensions["A"].width = 34
    g.column_dimensions["B"].width = 90
    lines = [("خلاصه", "")] + [(k, v) for k, v in summary.items()] + [
        ("", ""),
        ("روش", "برای هر بارکد به ترتیب: daychi (آدرس محصول = بارکد)، hamekara و rojashop (بارکد ثبت‌شده در صفحه)، "
                "Open Food Facts؛ اگر پیدا نشد جستجوی نام در دیجی‌کالا و ترب با کنترل شباهت نام و حجم/وزن."),
        ("اطمینان «بالا»", "منبع همین بارکد را دقیقاً به محصول وصل کرده و نام هم می‌خواند."),
        ("اطمینان «متوسط»", "بارکد دقیق ولی نام متفاوت، یا تطبیق قوی نام + حجم بدون بارکد. قبل از استفاده نگاه کنید."),
        ("اطمینان «پایین»", "فقط شباهت نسبی نام یا اختلاف حجم/وزن. حتماً دستی بررسی شود."),
        ("پیدا نشد", "هیچ منبعی عکس معتبر نداشت؛ عکس دستی لازم است. هیچ عکسی ساخته یا حدس زده نمی‌شود."),
        ("هشدار", "بسته‌بندی قدیم/جدید ممکن است با یک بارکد عکس متفاوت داشته باشد؛ بستهٔ چندتایی معمولاً بارکد جدا دارد."),
    ]
    for a, b in lines:
        g.append([a, b])
    for c in g["A"]:
        c.font = Font(bold=True)
    for row in g.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right")
    wb.save(out)
    shutil.rmtree(tmp, ignore_errors=True)
    return n_img


def main(argv=None):
    ap = argparse.ArgumentParser(description="بارکدها را از اکسل می‌خواند، عکس کالا را پیدا می‌کند و اکسل خروجی با عکس می‌سازد.")
    ap.add_argument("input", help="فایل اکسل ورودی (ستون بارکد لازم است)")
    ap.add_argument("-o", "--output", help="فایل اکسل خروجی (پیش‌فرض: <ورودی>-with-photos.xlsx)")
    ap.add_argument("--sheet", help="نام شیت (پیش‌فرض: اولین شیت)")
    ap.add_argument("--barcode-col", help="نام ستون بارکد (پیش‌فرض: «بارکد»)")
    ap.add_argument("--name-col", help="نام ستون نام کالا (پیش‌فرض: «نام محصول»/«نام کالا»)")
    ap.add_argument("--image-col", help="ستونی که آدرس عکس در آن نوشته شود (پیش‌فرض: «آدرس عکس اصلی»)")
    ap.add_argument("--sources", default=DEFAULT_SOURCES, help=f"منابع به ترتیب (پیش‌فرض: {DEFAULT_SOURCES})")
    ap.add_argument("--delay", type=float, default=1.5, help="حداقل فاصلهٔ بین درخواست‌ها به ثانیه (پیش‌فرض 1.5)")
    ap.add_argument("--min-score", type=float, default=0.6, help="حداقل شباهت نام برای جستجوی نامی (0..1، پیش‌فرض 0.6)")
    ap.add_argument("--start", type=int, default=1, help="از ردیف چندم شروع شود (۱ = اولین کالا)")
    ap.add_argument("--limit", type=int, help="حداکثر تعداد کالا در این اجرا")
    ap.add_argument("--retry-missing", action="store_true", help="کالاهای «پیدا نشد» قبلی دوباره جستجو شوند")
    ap.add_argument("--cutout", action="store_true", help="حذف پس‌زمینه با rembg (pip install \"rembg[cpu]\")")
    ap.add_argument("--cutout-model", default="isnet-general-use", help="مدل rembg (برای بسته‌های سفید: u2net)")
    ap.add_argument("--thumb", type=int, default=120, help="ارتفاع عکس داخل اکسل به پیکسل (پیش‌فرض 120)")
    ap.add_argument("-v", "--verbose", action="store_true")
    a = ap.parse_args(argv)

    sources = [s.strip() for s in a.sources.split(",") if s.strip()]
    bad = [s for s in sources if s not in SOURCES]
    if bad:
        sys.exit(f"منبع ناشناخته: {bad} — مجاز: {list(SOURCES)}")
    out = a.output or os.path.splitext(a.input)[0] + "-with-photos.xlsx"
    stem = os.path.splitext(out)[0]
    img_dir = stem + "_images"
    cut_dir = stem + "_cutouts"
    cache_path = stem + "_cache.json"
    os.makedirs(img_dir, exist_ok=True)

    headers, data, bi, ni, ci = read_input(a.input, a.sheet, a.barcode_col, a.name_col)
    print(f"{len(data)} کالا خوانده شد — ستون بارکد: «{headers[bi]}»، ستون نام: «{headers[ni] if ni is not None else '-'}»")

    cache: dict[str, dict] = {}
    if os.path.exists(cache_path):
        cache = json.load(open(cache_path, encoding="utf-8"))
        print(f"ادامهٔ اجرای قبلی: {len(cache)} کالا از قبل در {os.path.basename(cache_path)}")

    cutter = None
    if a.cutout:
        try:
            cutter = Cutter(a.cutout_model)
            os.makedirs(cut_dir, exist_ok=True)
        except ImportError:
            sys.exit('برای --cutout این را نصب کنید:  pip install "rembg[cpu]" scipy numpy')

    http = Http(a.delay, verbose=a.verbose)
    blocked: set[str] = set()
    results: list[Result | None] = [None] * len(data)
    todo = list(range(max(0, a.start - 1), len(data)))
    done_now = 0
    try:
        for idx in todo:
            row = data[idx]
            code = clean_barcode(row[bi]) or (clean_barcode(row[ci]) if ci is not None else "")
            name = str(row[ni]).strip() if ni is not None and row[ni] is not None else ""
            key = f"{code}|{name}"
            cached = cache.get(key)
            if cached and (cached["status"] != NOT_FOUND or not a.retry_missing) and \
                    (not cached["image_file"] or os.path.exists(cached["image_file"])):
                results[idx] = Result(**cached)
                continue
            if a.limit is not None and done_now >= a.limit:
                continue
            if len(blocked) == len(sources):
                print("همهٔ منابع مسدود شده‌اند؛ توقف. بعداً دوباره اجرا کنید تا ادامه دهد.")
                break
            done_now += 1
            print(f"[{idx + 1}/{len(data)}] {code} {name[:50]}", flush=True)
            r = lookup(http, code, name, sources, blocked, img_dir, a.min_score, a.verbose)
            if cutter and r.image_file:
                try:
                    cp = cutter(r.image_file, os.path.join(cut_dir, os.path.splitext(os.path.basename(r.image_file))[0] + ".png"))
                    if cp:
                        r.image_file = cp
                except Exception as e:
                    r.note += f"؛ حذف پس‌زمینه ناموفق ({type(e).__name__})"
            print(f"    → {r.status}  {r.source}  {r.source_title[:60]}", flush=True)
            results[idx] = r
            cache[key] = asdict(r)
            if done_now % 5 == 0:
                json.dump(cache, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, indent=0)
    except KeyboardInterrupt:
        print("\nمتوقف شد؛ نتایج تا اینجا ذخیره می‌شود. اجرای دوباره از همین‌جا ادامه می‌دهد.")
    json.dump(cache, open(cache_path, "w", encoding="utf-8"), ensure_ascii=False, indent=0)

    # fill rows that were done in earlier runs but outside this run's window
    for idx, row in enumerate(data):
        if results[idx] is None:
            code = clean_barcode(row[bi]) or (clean_barcode(row[ci]) if ci is not None else "")
            name = str(row[ni]).strip() if ni is not None and row[ni] is not None else ""
            c = cache.get(f"{code}|{name}")
            if c:
                results[idx] = Result(**c)

    checked = [r for r in results if r is not None]
    cnt = {k: sum(1 for r in checked if r.image_file and r.status == k) for k in (HIGH, MEDIUM, LOW)}
    summary = {
        "تعداد کل کالا": len(data),
        "بررسی‌شده": len(checked),
        "عکس پیدا شد": sum(cnt.values()),
        "اطمینان بالا / متوسط / پایین": f"{cnt[HIGH]} / {cnt[MEDIUM]} / {cnt[LOW]}",
        "پیدا نشد": sum(1 for r in checked if not r.image_file),
        "بررسی‌نشده": len(data) - len(checked),
        "منابع مسدودشده در این اجرا": "، ".join(sorted(blocked)) or "-",
        "حذف پس‌زمینه": "بله" if cutter else "خیر",
    }
    n_img = build_excel(out, headers, data, results, bi, a.image_col, a.thumb, summary)

    log_path = stem + "_log.csv"
    with open(log_path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["ردیف", "بارکد", "نام", "اطمینان", "منبع", "عنوان در منبع", "آدرس عکس", "لینک محصول",
                    "منابع بررسی‌شده", "یادداشت"])
        for idx, (row, r) in enumerate(zip(data, results), 1):
            if r:
                w.writerow([idx, r.barcode, row[ni] if ni is not None else "", r.status if r.image_file else NOT_FOUND,
                            r.source, r.source_title, r.image_url, r.page_url, " ".join(r.sources_tried), r.note])

    print("\n".join(f"  {k}: {v}" for k, v in summary.items()))
    print(f"خروجی: {out}  ({n_img} عکس داخل اکسل)\nعکس‌ها: {cutter and cut_dir or img_dir}\nلاگ: {log_path}")


if __name__ == "__main__":
    main()
